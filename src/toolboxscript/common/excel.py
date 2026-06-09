from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


def load_input_rows(path: str | Path) -> tuple[list[str], list[dict[str, str]]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(c.value or "").strip() for c in ws[1]]
    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {headers[i]: ("" if v is None else str(v)) for i, v in enumerate(row) if i < len(headers)}
        rows.append(item)
    return headers, rows


def copy_file(src: str | Path, dst: str | Path) -> None:
    from shutil import copy2
    copy2(src, dst)


def autosize_columns(ws, min_width: int = 10, max_width: int = 42):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            width = len(str(cell.value)) + 2
            widths[cell.column_letter] = max(widths.get(cell.column_letter, min_width), min(width, max_width))
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def style_header(ws):
    fill = PatternFill("solid", fgColor="D9E8FB")
    font = Font(bold=True)
    align = Alignment(vertical="center", wrap_text=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def style_all_cells(ws, wrap_columns: set[int] | None = None):
    wrap_columns = wrap_columns or set()
    for row in ws.iter_rows():
        for cell in row:
            if cell.column in wrap_columns:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_results_workbook(output_path: str | Path, summary_rows, all_rows, problem_rows, review_rows):
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for sheet_name, rows in [
        ("核验汇总", summary_rows),
        ("全量核验", all_rows),
        ("有问题企业", problem_rows),
        ("需复核企业", review_rows),
    ]:
        ws = wb.create_sheet(title=sheet_name)
        for r in rows:
            ws.append(list(r))
        style_header(ws)
        style_all_cells(ws, wrap_columns={1, 2, 3, 4, 5, 6, 7, 15, 16, 17, 18, 19})
        autosize_columns(ws)

    wb.save(output_path)

