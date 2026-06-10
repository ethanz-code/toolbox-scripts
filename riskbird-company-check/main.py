from __future__ import annotations

import argparse
import json
import os
import random
import runpy
import shutil
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BAD_STATUS_TERMS = ("注销", "吊销", "撤销", "迁出", "停业", "清算", "非正常")
BUSINESS_RATE_LIMIT_STATE = "operation:too:fast"
RESULT_COLUMNS = [
    "风鸟查询方式",
    "风鸟候选数量",
    "风鸟企业名称",
    "风鸟统一社会信用代码",
    "风鸟经营状态",
    "风鸟链接",
    "核验结果",
    "问题说明",
]


def now_stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d_%H%M%S")


def slugify(text: str) -> str:
    import re

    text = (text or "").strip().lower()
    text = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "-", text)
    text = re.sub(r"-+", "-", text).strip("-")
    return text or "run"


def ensure_dir(path: str | Path) -> Path:
    p = Path(path)
    p.mkdir(parents=True, exist_ok=True)
    return p


def new_run_dir(output_root: str | Path, script_name: str) -> Path:
    root = ensure_dir(output_root)
    base = root / f"{now_stamp()}_{slugify(script_name)}"
    run_dir = base
    idx = 2
    while run_dir.exists():
        run_dir = root / f"{base.name}-{idx}"
        idx += 1
    (run_dir / "raw" / "search").mkdir(parents=True, exist_ok=True)
    (run_dir / "output").mkdir(parents=True, exist_ok=True)
    (run_dir / "logs").mkdir(parents=True, exist_ok=True)
    return run_dir


def sleep_range(min_seconds: float, max_seconds: float) -> float:
    delay = random.uniform(min_seconds, max_seconds)
    time.sleep(delay)
    return delay


def backoff(attempt: int, base: float = 1.6, jitter: float = 0.35) -> float:
    delay = (base ** max(0, attempt - 1)) + random.uniform(0, jitter)
    time.sleep(delay)
    return delay


def business_rate_limit_message(payload) -> str:
    if not isinstance(payload, dict):
        return ""
    data = payload.get("data")
    if isinstance(data, dict) and data.get("state") == BUSINESS_RATE_LIMIT_STATE:
        return str(data.get("msg") or "操作过于频繁")
    return ""


def normalize_text(text: str | None) -> str:
    import re
    import unicodedata

    value = unicodedata.normalize("NFKC", text or "")
    return re.sub(r"\s+", "", value)


def normalize_code(text: str | None) -> str:
    return normalize_text(text).upper()


def safe_filename(text: str | None) -> str:
    import re

    value = normalize_text(text)
    value = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "_", value)
    value = re.sub(r"_+", "_", value).strip("_")
    return value or "unknown"


def load_input_rows(path: str | Path) -> tuple[list[str], list[dict[str, str]]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(c.value or "").strip() for c in ws[1]]
    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {headers[i]: ("" if v is None else str(v)) for i, v in enumerate(row) if i < len(headers)}
        rows.append(item)
    return headers, rows


def copy_file(src: str | Path, dst: str | Path) -> None:
    shutil.copy2(src, dst)


def _normalize_token(value: str | None) -> str:
    return (value or "").strip().strip('"').strip("'")


def _read_env_file(path: Path) -> str:
    # 支持 .env、单行 token 文件、或者简单的 key=value 配置。
    if not path.exists():
        return ""
    text = path.read_text(encoding="utf-8")
    if not text.strip():
        return ""
    if "\n" not in text and "=" not in text and " " not in text.strip():
        return _normalize_token(text)
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[len("export ") :].strip()
        if "=" not in line:
            continue
        key, value = line.split("=", 1)
        if key.strip() in {"RISKBIRD_TOKEN", "TOKEN", "token"}:
            return _normalize_token(value)
    return ""


def _read_py_file(path: Path) -> str:
    # 兼容 cookie.py / token.py 这类本地持久化会话文件。
    if not path.exists():
        return ""
    data = runpy.run_path(str(path))
    for key in ("RISKBIRD_TOKEN", "TOKEN", "token"):
        value = data.get(key)
        if value:
            return _normalize_token(str(value))
    return ""


def load_token(project_root: str | Path, token_env: str, token_file: str = "") -> str:
    # 优先读项目内配置，避免每次都手动传参。
    root = Path(project_root)
    if token_file:
        path = Path(token_file).expanduser()
        if path.suffix == ".py":
            return _read_py_file(path)
        return _read_env_file(path)

    env_path = root / ".env"
    token = _read_env_file(env_path)
    if token:
        return token

    py_path = root / "cookie.py"
    token = _read_py_file(py_path)
    if token:
        return token

    token = _normalize_token(os.environ.get(token_env) or os.environ.get("RISKBIRD_COOKIE"))
    if token:
        return token

    raise RuntimeError(f"missing token: set {token_env}, create .env, create cookie.py, or pass --token-file")


class JsonlLogger:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        ensure_dir(self.path.parent)

    def emit(self, event: str, **fields):
        payload = {"ts": datetime.now().isoformat(timespec="seconds"), "event": event, **fields}
        with self.path.open("a", encoding="utf-8") as f:
            f.write(json.dumps(payload, ensure_ascii=False) + "\n")
        return payload


class RiskbirdClient:
    def __init__(self, token: str):
        self.token = token
        self.base = "https://www.riskbird.com"

    def _headers(self) -> dict[str, str]:
        return {
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/125 Safari/537.36",
            "Cookie": f"token={self.token}" if self.token else "",
            "app-device": "WEB",
            "Referer": "https://www.riskbird.com/",
            "Origin": "https://www.riskbird.com",
            "Accept": "application/json",
            "Content-Type": "application/json",
        }

    def _request_json(self, method: str, url: str, *, body=None, params=None, retries: int = 4):
        # 所有请求统一走这里：补头、计时、重试、截断 401/429。
        last_error = None
        for attempt in range(1, retries + 1):
            full_url = url
            if params:
                query = urllib.parse.urlencode(params)
                full_url = f"{url}{'&' if '?' in url else '?'}{query}"
            data = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
            req = urllib.request.Request(full_url, data=data, headers=self._headers(), method=method)
            started = time.perf_counter()
            try:
                with urllib.request.urlopen(req, timeout=40) as resp:
                    text = resp.read().decode("utf-8", "ignore")
                    elapsed_ms = int((time.perf_counter() - started) * 1000)
                    try:
                        payload = json.loads(text) if text else None
                    except json.JSONDecodeError:
                        payload = None
                    rate_limit_msg = business_rate_limit_message(payload)
                    if rate_limit_msg:
                        last_error = RuntimeError(rate_limit_msg)
                        if attempt < retries:
                            backoff(attempt, base=4.0, jitter=1.5)
                            continue
                        return {
                            "ok": False,
                            "status": BUSINESS_RATE_LIMIT_STATE,
                            "url": full_url,
                            "elapsed_ms": elapsed_ms,
                            "text": text,
                            "data": payload,
                            "error": rate_limit_msg,
                        }
                    return {
                        "ok": True,
                        "status": resp.status,
                        "url": full_url,
                        "elapsed_ms": elapsed_ms,
                        "text": text,
                        "data": payload,
                    }
            except urllib.error.HTTPError as exc:
                text = exc.read().decode("utf-8", "ignore") if hasattr(exc, "read") else ""
                elapsed_ms = int((time.perf_counter() - started) * 1000)
                if exc.code in {401, 429}:
                    return {
                        "ok": False,
                        "status": exc.code,
                        "url": full_url,
                        "elapsed_ms": elapsed_ms,
                        "text": text,
                        "error": str(exc),
                    }
                last_error = exc
                if attempt < retries:
                    backoff(attempt)
                    continue
                return {
                    "ok": False,
                    "status": exc.code,
                    "url": full_url,
                    "elapsed_ms": elapsed_ms,
                    "text": text,
                    "error": str(exc),
                }
            except Exception as exc:
                last_error = exc
                if attempt < retries:
                    backoff(attempt)
                    continue
        return {
            "ok": False,
            "status": None,
            "url": url,
            "elapsed_ms": None,
            "text": "",
            "error": str(last_error) if last_error else "unknown error",
        }

    def company_search(self, key: str):
        return self._request_json(
            "POST",
            f"{self.base}/riskbird-api/api/v1/companies/search",
            body={
                "queryType": "1",
                "searchKey": key,
                "pageNo": 1,
                "range": 10,
                "selectConditionData": json.dumps({"status": "", "sort_field": "", "keyword_fields": ""}, ensure_ascii=False),
            },
        )


def extract_candidates(payload) -> list[dict]:
    if isinstance(payload, dict):
        for key in ("list", "records", "data"):
            value = payload.get(key)
            if isinstance(value, list):
                return [item for item in value if isinstance(item, dict)]
        if isinstance(payload.get("data"), dict):
            return extract_candidates(payload["data"])
    if isinstance(payload, list):
        return [item for item in payload if isinstance(item, dict)]
    return []


def candidate_name(item: dict) -> str:
    return normalize_text(item.get("entName") or item.get("ENTNAME") or item.get("name") or item.get("perName") or "")


def candidate_code(item: dict) -> str:
    for key in ("uniscid", "taxId", "creditCode", "code"):
        value = item.get(key)
        if value:
            return normalize_code(value)
    return ""


def candidate_score(item: dict, wanted_name: str, wanted_code: str) -> int:
    score = 0
    name = candidate_name(item)
    code = candidate_code(item)
    wanted_name_norm = normalize_text(wanted_name)
    wanted_code_norm = normalize_code(wanted_code)
    if wanted_code_norm and code and code == wanted_code_norm:
        score += 120
    if wanted_name and item.get("entName") == wanted_name:
        score += 80
    if wanted_name_norm and name == wanted_name_norm:
        score += 70
    elif wanted_name_norm and (wanted_name_norm in name or name in wanted_name_norm):
        score += 45
    if item.get("entStatus") or item.get("ENTSTATUS"):
        score += 3
    if wanted_code_norm and code and code != wanted_code_norm:
        score -= 15
    return score


def choose_best_candidate(candidates: list[dict], wanted_name: str, wanted_code: str):
    scored = [(candidate_score(item, wanted_name, wanted_code), item) for item in candidates]
    if not scored:
        return None, 0
    scored.sort(key=lambda x: x[0], reverse=True)
    return scored[0][1], scored[0][0]


def company_status(item: dict) -> str:
    return (item.get("entStatus") or item.get("ENTSTATUS") or item.get("status") or item.get("entStatusText") or "").strip()


def status_is_problem(status: str) -> bool:
    return any(term in (status or "") for term in BAD_STATUS_TERMS)


def build_company_link(name: str, entid: str) -> str:
    if not entid:
        return ""
    return f"https://www.riskbird.com/ent/{urllib.parse.quote(name)}.html?entid={entid}"


def search_best_company(client: RiskbirdClient, name: str, code: str):
    # 先按统一社会信用代码查，再按名称查，尽量少打请求。
    attempts = []
    candidates: list[dict] = []
    blocked = None

    def run_query(strategy: str, key: str):
        nonlocal blocked
        if not key:
            return
        response = client.company_search(key)
        items = extract_candidates(response.get("data") if response.get("ok") else None)
        attempts.append(
            {
                "strategy": strategy,
                "key": key,
                "response": response,
                "items": items,
            }
        )
        if not response.get("ok") and response.get("status") in {401, 429, BUSINESS_RATE_LIMIT_STATE}:
            blocked = response
        candidates.extend(items)

    if code:
        run_query("code", code)
        best, score = choose_best_candidate(candidates, name, code)
        if blocked or (best and candidate_code(best) == normalize_code(code)):
            return attempts, candidates, best, score, blocked

    if name:
        if attempts:
            sleep_range(2.5, 4.5)
        run_query("name", name)

    best, score = choose_best_candidate(candidates, name, code)
    return attempts, candidates, best, score, blocked


def classify_row(row: dict, best: dict | None, candidates: list[dict], query_path: str, blocked: dict | None):
    # 只依赖 search 返回里的经营状态，不再请求详情页。
    name = (row.get("企业名称") or "").strip()
    code = (row.get("统一社会信用代码") or "").strip()

    if blocked:
        return {
            **row,
            "风鸟查询方式": query_path,
            "风鸟候选数量": len(candidates),
            "风鸟企业名称": (best or {}).get("entName") or (best or {}).get("ENTNAME") or "",
            "风鸟统一社会信用代码": candidate_code(best or {}) if best else "",
            "风鸟经营状态": company_status(best or {}) if best else "",
            "风鸟链接": build_company_link((best or {}).get("entName") or name, (best or {}).get("entid") or ""),
            "核验结果": "查询受限",
            "问题说明": f"风鸟搜索接口返回 {blocked.get('status')}，已停止继续查询",
        }

    if not candidates:
        return {
            **row,
            "风鸟查询方式": query_path,
            "风鸟候选数量": 0,
            "风鸟企业名称": "",
            "风鸟统一社会信用代码": "",
            "风鸟经营状态": "",
            "风鸟链接": "",
            "核验结果": "未查到",
            "问题说明": "搜索接口未返回候选结果",
        }

    hit_name = (best or {}).get("entName") or (best or {}).get("ENTNAME") or ""
    hit_code = candidate_code(best or {})
    status = company_status(best or {})
    issues = []
    if hit_code and normalize_code(code) and hit_code != normalize_code(code):
        issues.append(f"信用代码不一致：表格为{code}，风鸟返回{hit_code}")
    if not status:
        issues.append("未返回经营状态")

    if status and status_is_problem(status):
        result = "有问题"
        issues.insert(0, f"经营状态：{status}")
    elif status:
        result = "正常经营"
    else:
        result = "未查到"

    return {
        **row,
        "风鸟查询方式": query_path,
        "风鸟候选数量": len(candidates),
        "风鸟企业名称": hit_name,
        "风鸟统一社会信用代码": hit_code,
        "风鸟经营状态": status,
        "风鸟链接": build_company_link(hit_name or name, (best or {}).get("entid") or ""),
        "核验结果": result,
        "问题说明": "；".join(issues),
    }


class LiveResultsWorkbook:
    def __init__(self, output_path: str | Path, headers: list[str], input_path: str = "", run_dir: str = ""):
        self.output_path = Path(output_path)
        ensure_dir(self.output_path.parent)
        self.headers = list(headers)
        self.input_path = input_path
        self.run_dir = run_dir
        self.total = 0
        self.counts = {"正常经营": 0, "有问题": 0, "未查到": 0, "查询受限": 0}

        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.ws_summary = self.wb.create_sheet("核验汇总")
        self.ws_all = self.wb.create_sheet("全量核验")
        self.ws_problem = self.wb.create_sheet("有问题企业")

        self._init_summary_sheet()
        self._init_data_sheet(self.ws_all)
        self._init_data_sheet(self.ws_problem)
        self.save()

    def _init_data_sheet(self, ws):
        ws.append(self.headers)
        fill = PatternFill("solid", fgColor="D9E8FB")
        font = Font(bold=True)
        align = Alignment(vertical="center", wrap_text=True)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = align
        ws.freeze_panes = "A2"

    def _init_summary_sheet(self):
        ws = self.ws_summary
        rows = [
            ["项目", "数量"],
            ["总行数", 0],
            ["正常经营", 0],
            ["有问题", 0],
            ["未查到", 0],
            ["查询受限", 0],
            ["", ""],
            ["输入文件", self.input_path],
            ["运行目录", self.run_dir],
        ]
        for row in rows:
            ws.append(row)
        fill = PatternFill("solid", fgColor="D9E8FB")
        font = Font(bold=True)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
        ws.freeze_panes = "A2"

    def _extract_values(self, row: dict) -> list:
        return [row.get(header, "") for header in self.headers]

    def _append_row(self, ws, values):
        ws.append(values)
        row_idx = ws.max_row
        for cell in ws[row_idx]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    def _update_summary(self):
        ws = self.ws_summary
        ws["B2"] = self.total
        ws["B3"] = self.counts["正常经营"]
        ws["B4"] = self.counts["有问题"]
        ws["B5"] = self.counts["未查到"]
        ws["B6"] = self.counts["查询受限"]

    def append_result(self, row: dict):
        # 每处理一条就落盘，防止中途断掉时丢结果。
        result = row.get("核验结果", "")
        self.total += 1
        if result in self.counts:
            self.counts[result] += 1

        values = self._extract_values(row)
        self._append_row(self.ws_all, values)
        if result == "有问题":
            self._append_row(self.ws_problem, values)
        self._update_summary()
        self.save()

    def finalize(self):
        for ws in (self.ws_summary, self.ws_all, self.ws_problem):
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                max_width = 10
                letter = col[0].column_letter
                for cell in col:
                    if cell.value is None:
                        continue
                    max_width = max(max_width, min(len(str(cell.value)) + 2, 42))
                ws.column_dimensions[letter].width = max_width
        self.save()

    def save(self):
        self.wb.save(self.output_path)


def run_riskbird_company_check(args):
    input_path = Path(args.input).expanduser().resolve()
    project_root = Path(__file__).resolve().parent
    run_dir = new_run_dir(args.output_root, "riskbird-company-check")
    logger = JsonlLogger(run_dir / "logs" / "log.jsonl")
    logger.emit("run.start", input=str(input_path), run_dir=str(run_dir))

    if args.save_input_copy:
        copy_file(input_path, run_dir / "input.xlsx")

    token = load_token(project_root, args.token_env, args.token_file)
    client = RiskbirdClient(token)

    input_headers, rows = load_input_rows(input_path)
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    output_headers = input_headers + [col for col in RESULT_COLUMNS if col not in input_headers]
    workbook_path = run_dir / "output" / "企业经营核验结果.xlsx"
    workbook = LiveResultsWorkbook(workbook_path, output_headers, input_path=str(input_path), run_dir=str(run_dir))

    results = []
    for idx, row in enumerate(rows, 1):
        name = (row.get("企业名称") or "").strip()
        code = (row.get("统一社会信用代码") or "").strip()
        row_id = safe_filename(code or name or str(idx))
        logger.emit("item.start", idx=idx, name=name, code=code)

        attempts, candidates, best, score, blocked = search_best_company(client, name, code)
        (run_dir / "raw" / "search" / f"{row_id}.json").write_text(
            json.dumps(attempts, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        query_path = " | ".join(f"{item['strategy']}({len(item['items'])})" for item in attempts)
        result_row = classify_row(row, best, candidates, query_path, blocked)
        results.append(result_row)
        workbook.append_result(result_row)

        logger.emit(
            "item.done",
            idx=idx,
            name=name,
            code=code,
            result=result_row["核验结果"],
            candidates=len(candidates),
            score=score,
            status=result_row["风鸟经营状态"],
        )

        if blocked:
            logger.emit("run.blocked", idx=idx, status=blocked.get("status"), message="search endpoint blocked")
            break

        sleep_range(args.sleep_min, args.sleep_max)

    workbook.finalize()
    summary = {
        "total": len(results),
        "正常经营": sum(1 for r in results if r["核验结果"] == "正常经营"),
        "有问题": sum(1 for r in results if r["核验结果"] == "有问题"),
        "未查到": sum(1 for r in results if r["核验结果"] == "未查到"),
        "查询受限": sum(1 for r in results if r["核验结果"] == "查询受限"),
    }
    (run_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (run_dir / "output" / "results.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    logger.emit("run.finish", **summary, workbook=str(workbook_path))
    print(json.dumps({"run_dir": str(run_dir), "workbook": str(workbook_path), **summary}, ensure_ascii=False))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="riskbird-company-check")
    parser.add_argument("--input", required=True, help="输入 xlsx 文件路径")
    parser.add_argument("--output-root", default="runs", help="运行输出根目录")
    parser.add_argument("--limit", type=int, default=0, help="只处理前 N 行，用于测试")
    parser.add_argument("--sleep-min", type=float, default=4.0, help="单条企业最小间隔秒数")
    parser.add_argument("--sleep-max", type=float, default=8.0, help="单条企业最大间隔秒数")
    parser.add_argument("--token-env", "--cookie-env", dest="token_env", default="RISKBIRD_TOKEN", help="风鸟 token 环境变量名")
    parser.add_argument("--token-file", "--cookie-file", dest="token_file", default="", help="可选，风鸟 token 文件路径")
    parser.add_argument("--save-input-copy", action="store_true", help="把输入文件拷贝到运行目录")
    parser.set_defaults(func=run_riskbird_company_check)
    return parser


def main(argv: list[str] | None = None) -> None:
    parser = build_parser()
    args = parser.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main(sys.argv[1:])
